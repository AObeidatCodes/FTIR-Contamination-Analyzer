"""
FTIR Contamination Analyzer - the records, formatted for Excel

The records themselves (sample log, audit log, signatures, library change log) stay plain
CSV files that are only ever added to: that is what makes them trustworthy, and Excel
would lock them while they are open. To read them, the Records menu makes a formatted copy
(an .xlsx "view") and opens that:

    a title, what the record is and where the original file is (with its SHA-256)
    a header row that stays in place while scrolling, with filters
    dates as real dates, numbers as numbers, verdicts and actions in their colours
    long text wrapped, checksums in a fixed-width font, ready to print (landscape, one page wide)
    a second sheet that explains every column

The view is made fresh every time, in a temporary folder: changing it changes nothing.
The report's results table is written the same way (report.xlsx).
Needs the 'openpyxl' package (in requirements.txt).

Short names used in this file:
    ws = worksheet    wb = workbook    col = column number (1 = A)    t = theme colours

Author : Ali Obeidat
Version: 1.0.0
License: MIT (see LICENSE)
"""

import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __author__, __title__, __version__
from .records import DECISIONS, readable_time, who_text
from .spectrum_io import checksum, read_table
from .theme import LIGHT

t = {k: v.lstrip("#") for k, v in LIGHT.items() if isinstance(v, str) and v.startswith("#")}
TEXT_FONT, MONO_FONT = "Aptos", "Consolas"      # Excel's own font (falls back to Calibri), and one for checksums
FIRST_ROW = 5                                   # the header row; title, explanation and source above it

# What each column means (shown on the second sheet)
COLUMN_HELP = {
    "file": "the raw spectrum file", "sha256": "SHA-256 checksum of that exact file",
    "mode": "how it was measured (ATR or transmission)", "recorded_by": "login name of the person",
    "recorded_at": "the date it was recorded", "note": "how the answer was given",
    "run_at": "when the report was made", "run_by": "login name of the person who made it",
    "program": "program and version", "code_sha256": "checksum of the program's code (first 16 characters)",
    "library_sha256": "checksum of the reference library (first 16 characters)",
    "file_sha256": "SHA-256 checksum of the file", "sample": "the sample (spectrum) name",
    "verdict": "probable, weak or no reliable match", "match": "the material the verdict names",
    "r": "the score (Pearson r of the derivative fingerprints)", "flags": "things to look at",
    "time": "when the change was made", "user": "login name of the person", "action": "add or remove",
    "material": "the reference", "category": "where it comes from in a fab", "fab_sources": "where it comes from",
    "check_next": "what to check next", "confirmation": "how its identity was confirmed", "notes": "notes",
    "source_file": "the spectrum file it was made from", "source_sha256": "checksum of that file",
    "reason": "why it was removed", "signed_at": "when it was signed", "signed_by": "the name of the person",
    "login": "their login name", "decision": "accepted, accepted with changes, or re-measure",
    "comment": "the reviewer's comment", "report": "the report file that was signed",
    "report_sha256": "checksum of that report file", "spectra": "how many spectra the report has",
    "unit": "transmittance (%T) or absorbance (A) in the file", "mode_source": "where the mode came from",
    "compared_as": "the mode the comparison used", "best": "the best-scoring reference", "gap":
    "lead over the next different material", "second": "the next different material", "ch": "C-H present?",
    "ch_share": "C-H band height / tallest band", "oh": "O-H / N-H", "oh_share": "O-H band height / tallest band",
    "max_A": "highest absorbance", "co2_share": "CO2 band height / tallest band", "noise": "noise (absorbance)",
    "key_bands_found": "key bands of the match found", "key_bands": "key bands of the match",
    "unexplained": "bands the match cannot explain", "next_step": "what to do next",
}
WIDE = {"note", "flags", "fab_sources", "check_next", "confirmation", "notes", "reason", "comment", "next_step",
        "report", "unexplained"}
# What each record is, for records that do not explain themselves in their first lines
ABOUT = {
    "Audit log": "One row per analysed spectrum in every report that was made: which file (and its SHA-256), the "
                 "result, when, who, which program and which library. Rows are only ever added.",
    "Sample log": "How each raw file was measured (ATR or transmission), recorded once per exact file (identified "
                  "by its SHA-256), never inside the raw file itself. Rows are only ever added.",
    "Signatures": "One row per signed report: who signed, when, the decision and the comment, and the SHA-256 of "
                  "the report file that was signed. Rows are only ever added.",
    "Library change log": "Every addition to and removal from the in-house library, in order, with who, when, why "
                          "and the checksums of the files. Rows are only ever added; nothing is deleted.",
}
HEADINGS = {"r": "Score (r)", "gap": "Lead", "max_A": "Highest A", "ch": "C-H", "oh": "O-H / N-H",
            "ch_share": "C-H share", "oh_share": "O-H share", "co2_share": "CO2 share"}
TIMES = {"run_at", "recorded_at", "time", "signed_at"}
VERDICT_COLOURS = {"probable": ("good", "good_bg"), "weak": ("warn", "warn_bg"),
                   "no reliable match": ("none", "none_bg"), "add": ("good", "good_bg"), "remove": ("bad", "bad_bg"),
                   "accept": ("good", "good_bg"), "accept with changes": ("warn", "warn_bg"),
                   "re-measure": ("bad", "bad_bg")}


def heading(name):
    """'file_sha256' -> 'File SHA-256', 'r' -> 'Score (r)'."""
    if name in HEADINGS:
        return HEADINGS[name]
    words = str(name).replace("sha256", "SHA-256").split("_")
    text = " ".join(words)
    return text[:1].upper() + text[1:]


def _as_time(text):
    """An ISO time from a record -> a datetime Excel shows as a date (wall-clock time, as recorded)."""
    try:
        moment = datetime.fromisoformat(str(text))
    except ValueError:
        return None
    return moment.replace(tzinfo=None)


def _as_number(text):
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def _font(colour="ink", size=11, bold=False, mono=False):
    return Font(name=MONO_FONT if mono else TEXT_FONT, size=size, bold=bold, color=t[colour])


def _table(ws, frame, top=FIRST_ROW):
    """The table, from row `top` down: styled header, typed cells, widths, filters, frozen header."""
    line = Side(style="thin", color=t["line"])
    header_fill = PatternFill("solid", fgColor=t["surface_2"])
    columns = list(frame.columns)
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=top, column=col, value=heading(name))
        cell.font = _font("ink_2", 11, bold=True)
        cell.fill = header_fill
        cell.border = Border(bottom=Side(style="medium", color=t["line_strong"]))
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[top].height = 30
    for row, values in enumerate(frame.itertuples(index=False), start=top + 1):
        for col, (name, value) in enumerate(zip(columns, values), start=1):
            text = "" if value is None or (isinstance(value, float) and pd.isna(value)) else value
            cell = ws.cell(row=row, column=col)
            cell.border = Border(bottom=line)
            cell.alignment = Alignment(vertical="top", wrap_text=name in WIDE)
            cell.font = _font()
            moment = _as_time(text) if name in TIMES and text != "" else None
            number = _as_number(text) if name in ("r", "gap", "ch_share", "oh_share", "max_A", "co2_share",
                                                   "noise", "spectra", "key_bands_found", "key_bands") else None
            if moment is not None:
                cell.value = moment
                cell.number_format = "d mmm yyyy" if len(str(text)) <= 10 else "d mmm yyyy, hh:mm"
                cell.alignment = Alignment(vertical="top", horizontal="left")
            elif number is not None and text != "":
                cell.value = number
                cell.number_format = "0" if name in ("spectra", "key_bands_found", "key_bands") else (
                    "0.0000" if name == "noise" else "0.00")
            else:
                cell.value = str(text)
            if "sha256" in name:
                if len(str(text)) > 24:                                 # the start is enough to recognise it;
                    cell.value = f"{str(text)[:16]}…"                   # the full value is in the record
                cell.font = _font("muted", 10, mono=True)
            colours = VERDICT_COLOURS.get(str(text)) if name in ("verdict", "action", "decision") else None
            if colours:
                cell.font = _font(colours[0], 11, bold=True)
                cell.fill = PatternFill("solid", fgColor=t[colours[1]])
                if name == "decision":
                    cell.value = DECISIONS.get(str(text), str(text))
    for col, name in enumerate(columns, start=1):                       # widths from the content
        longest = max([len(heading(name))] + [len(str(v)) for v in frame[name].head(500)])
        if "sha256" in name:
            width = 20
        elif name in WIDE:
            width = min(max(longest, 18), 48)
        elif name in TIMES:
            width = 19
        else:
            width = min(max(longest + 2, 8), 36)
        ws.column_dimensions[get_column_letter(col)].width = width
    last = get_column_letter(max(1, len(columns)))
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    ws.auto_filter.ref = f"A{top}:{last}{top + max(1, len(frame))}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = f"{top}:{top}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _title(ws, title, explanation, source, width):
    last = get_column_letter(max(width, 6))
    ws["A1"] = title
    ws["A1"].font = _font("ink", 18, bold=True)
    ws.row_dimensions[1].height = 30
    for row, text, colour in ((2, explanation, "ink_2"), (3, source, "muted")):
        ws.cell(row=row, column=1, value=text).font = _font(colour, 11 if row == 2 else 10)
        ws.merge_cells(f"A{row}:{last}{row}")
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 32 if row == 2 else 28


def _about(wb, title, columns, facts):
    """A second sheet: what the record is, and what each column means."""
    ws = wb.create_sheet("About")
    ws["A1"] = f"About this {title.lower()}"
    ws["A1"].font = _font("ink", 16, bold=True)
    row = 3
    for key, value in facts:
        ws.cell(row=row, column=1, value=key).font = _font("muted")
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = _font("ink", mono=key.endswith("SHA-256"))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Column").font = _font("ink_2", bold=True)
    ws.cell(row=row, column=2, value="What it means").font = _font("ink_2", bold=True)
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=t["surface_2"])
    for name in columns:
        row += 1
        ws.cell(row=row, column=1, value=heading(name)).font = _font()
        ws.cell(row=row, column=2, value=COLUMN_HELP.get(name, "")).font = _font("ink_2")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    ws.sheet_view.showGridLines = False


def _explanation(path):
    """The record's own explanation: its '#' lines at the top, as one sentence."""
    lines = []
    for line in Path(path).read_text(encoding="utf-8", errors="replace").splitlines():
        if not line.startswith("#"):
            break
        lines.append(line.lstrip("#").strip())
    return " ".join(lines)


def write_log_workbook(csv_path, out_path, title):
    """A formatted copy of one record (a CSV log) as an Excel workbook."""
    csv_path = Path(csv_path)
    frame = read_table(csv_path, dtype=str).fillna("")
    made = datetime.now().astimezone().isoformat(timespec="seconds")
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    source = (f"Record: {csv_path}  ·  {len(frame)} row{'s' if len(frame) != 1 else ''}  ·  this sheet is a copy "
              f"made {readable_time(made)}; the record itself is the CSV file, which is only ever added to.")
    explanation = ABOUT.get(title) or _explanation(csv_path)
    _title(ws, title, explanation, source, len(frame.columns))
    _table(ws, frame)
    _about(wb, title, list(frame.columns), [
        ("Record", str(csv_path)), ("Record SHA-256", checksum(csv_path)), ("Rows", str(len(frame))),
        ("This copy made", readable_time(made, with_zone=True)), ("Made with", f"{__title__} {__version__}"),
        ("What it is", explanation),
        ("Checksums", "shortened here to their first 16 characters; the full values are in the record")])
    wb.properties.creator = __author__
    wb.properties.title = title
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def write_results_workbook(table, info, out_path, signoff=None):
    """report.xlsx: the results of one report, formatted, with who, when and the sign-off above them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    n = len(table)
    analysed_by = who_text(info.get("run_by_name", ""), info["run_by"])
    signed = (f"Signed by {who_text(signoff['name'], signoff['login'])} on {readable_time(signoff['signed_at'])}: "
              f"{DECISIONS[signoff['decision']]}." if signoff else "Not signed in the program.")
    _title(ws, "FTIR contamination report",
           f"{n} spectr{'um' if n == 1 else 'a'}, analysed {readable_time(info['run_at'], with_zone=True)} by "
           f"{analysed_by}. {signed}",
           f"{info['library']}  ·  library SHA-256 {info['library_sha256'][:16]}…  ·  {info['program']}",
           len(table.columns))
    _table(ws, table.astype(object).where(table.notna(), ""))
    facts = [("Analysed", readable_time(info["run_at"], with_zone=True)), ("By", analysed_by),
             ("Program", info["program"]), ("Code SHA-256", info["code_sha256"]), ("Software", info["software"]),
             ("Library", info["library"]), ("Library SHA-256", info["library_sha256"]),
             ("Impostor test", f"best wrong answer {info['validation'].score:.2f} "
                               f"({info['validation'].material} -> {info['validation'].impostor})"),
             ("Sign-off", signed + (f" Comment: {signoff['comment']}" if signoff and signoff.get("comment") else ""))]
    _about(wb, "report", list(table.columns), facts)
    wb.properties.creator = __author__
    wb.properties.title = "FTIR contamination report"
    wb.save(out_path)
    return Path(out_path)


def view_of(csv_path, title, folder=None):
    """Make a fresh formatted copy of a record in a temporary folder and return its path."""
    folder = Path(folder or Path(tempfile.gettempdir()) / "FTIR Contamination Analyzer" / "views")
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    return write_log_workbook(csv_path, folder / f"{title} {stamp}.xlsx", title)
