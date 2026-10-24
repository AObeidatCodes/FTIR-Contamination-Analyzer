"""
FTIR Contamination Analyzer - tests

Known-answer tests: every demo spectrum must give the verdict it was validated with,
the raw data must never change, and the records must only ever grow. ATR and transmission
meet in the library's look whichever side was measured which way; an unknown mode is worked
out from the better fit, and a stated mode that fits clearly worse is flagged.

    python -m pytest            (from the project folder)

Author : Ali Obeidat
Version: 1.1.0
License: MIT (see LICENSE)
"""

import re
from dataclasses import replace

import numpy as np
import pandas as pd
import pytest

from ftir_analyzer import config
from ftir_analyzer.analysis import NO_MATCH, PROBABLE, WEAK, analyse, decide
from ftir_analyzer.checks import describe_vapour
from ftir_analyzer.library import impostor_test, load_library, pearson
from ftir_analyzer.processing import GRID, fingerprint, on_grid, to_library_mode, transmittance_to_absorbance
from ftir_analyzer.records import SampleLog, resolve_modes, run_info
from ftir_analyzer.report import write_report
from ftir_analyzer.spectrum_io import SpectrumFileError, read_spectra

EXAMPLES = config.EXAMPLES_DIR


@pytest.fixture(scope="module")
def library():
    """The shipped library, loaded once for all tests in this file."""
    return load_library()


def example(name, library, mode=None):
    """Analyse one demo file; the mode comes from the file note unless one is given."""
    spectrum = read_spectra(EXAMPLES / name)[0]
    return analyse(spectrum, library, mode or "unknown")


# ---------------------------------------------------------------- reading files
def test_raw_data_is_read_only():
    spectrum = read_spectra(EXAMPLES / "01_wipe_area_particle_ATR.csv")[0]
    with pytest.raises(ValueError):                    # numpy refuses to write into a read-only array
        spectrum.raw[0] = 0.0


def test_headerless_tab_separated_file(tmp_path):
    wn = np.arange(4000, 600, -10.0)                   # high -> low, as many instruments export
    a = np.exp(-((wn - 1700) / 20) ** 2)
    path = tmp_path / "export.txt"
    path.write_text("\n".join(f"{x}\t{y}" for x, y in zip(wn, a)))
    spectrum = read_spectra(path)[0]
    assert spectrum.unit == "A"
    assert spectrum.wn[0] < spectrum.wn[-1]            # sorted low -> high
    assert spectrum.raw[np.argmin(np.abs(spectrum.wn - 1700))] == pytest.approx(1.0)   # pairs stayed together


def test_transmittance_recognised_from_values(tmp_path):
    path = tmp_path / "film.csv"
    wn = np.arange(650, 4000, 5.0)
    pd.DataFrame({"wavenumber": wn, "signal": 90 - 40 * np.exp(-((wn - 1450) / 15) ** 2)}).to_csv(path, index=False)
    assert read_spectra(path)[0].unit == "%T"


@pytest.mark.parametrize("text", ["x,y\n1,2\n2,3\n", "\n".join(f"{i},{i}" for i in range(10000, 10100))])
def test_unusable_files_give_a_clear_error(tmp_path, text):
    path = tmp_path / "bad.csv"
    path.write_text(text)
    with pytest.raises(SpectrumFileError):
        read_spectra(path)


# ---------------------------------------------------------------- processing
def test_absorbance_from_transmittance():
    assert transmittance_to_absorbance([100, 10, 1, 0]) == pytest.approx([0, 1, 2, 6])   # 0 % T is clipped at A = 6


def test_grid_never_invents_data():
    y = on_grid(np.array([1000.0, 2000.0]), np.array([1.0, 1.0]))
    assert np.isnan(y[GRID < 1000]).all() and np.isnan(y[GRID > 2000]).all()


def test_atr_correction_returns_a_new_array():
    wn, a = np.array([1000.0, 2000.0]), np.array([1.0, 1.0])
    out = to_library_mode(wn, a, "transmission", "ATR")
    assert out == pytest.approx([1.0, 0.5])            # factor 1000 / wavenumber
    assert out is not a and a.tolist() == [1.0, 1.0]   # the input is untouched
    assert to_library_mode(wn, a, "unknown", "ATR") is not a


def test_fingerprint_ignores_offset_and_thickness():
    y = np.exp(-((GRID - 1700) / 15) ** 2) + np.exp(-((GRID - 2900) / 30) ** 2)
    assert pearson(fingerprint(y), fingerprint(3 * y + 0.5)) == pytest.approx(1.0)


# ---------------------------------------------------------------- verdict rule
def test_verdict_rule():
    assert decide([("PET", 0.97), ("PBT", 0.87)]) == (PROBABLE, "PET")
    assert decide([("PET", 0.97), ("PBT", 0.93)]) == (WEAK, "PET or PBT")        # too close to call
    assert decide([("PS", 0.80), ("PBT", 0.35)]) == (WEAK, "PS")
    assert decide([("PS", 0.27), ("PBT", 0.01)]) == (NO_MATCH, "-")


def test_impostor_test_passes_on_the_shipped_library(library):
    validation = impostor_test(library)
    assert validation.passed and validation.score < config.PROBABLE_R


# ---------------------------------------------------------------- known answers
@pytest.mark.parametrize("name, verdict, match", [
    ("01_wipe_area_particle_ATR.csv", PROBABLE, "PET"),
    ("02_wafer_box_particle_ATR.csv", PROBABLE, "PE (high density)"),
    ("03_litho_rework_film_ATR.csv", PROBABLE, "PMMA"),
    ("04_post_CMP_clean_particle_ATR.csv", PROBABLE, "PVOH"),
    ("05_fibre_on_wafer_ATR.csv", WEAK, "Nylon 6,6"),
    ("06_PS_calibration_film_transmission.csv", WEAK, "PS"),
    ("08_fibre_on_filter.csv", NO_MATCH, "-"),
    ("10_tool_cover_fragment_ATR.csv", WEAK, "PS"),                 # ABS is not in the shipped library
    ("11_second_fibre_on_filter.csv", NO_MATCH, "-"),                # nor is cellulose acetate
    ("12_shipping_box_film_ATR.csv", PROBABLE, "EVOH (44% ethylene)"),
])
def test_demo_spectra(library, name, verdict, match):
    mode = "transmission" if "transmission" in name else "ATR" if "ATR" in name else None
    result = example(name, library, mode)
    assert (result.verdict, result.match) == (verdict, match)


def test_second_component_is_found(library):
    result = example("07_particle_on_filter.csv", library, "ATR")
    assert result.verdict == PROBABLE and result.match == "PE (high density)"
    assert any(abs(w - 1030) <= 15 for w in result.evidence.unexplained)          # Si-O from the soil


def test_empty_spot_is_recognised(library):
    result = example("09_empty_spot_on_filter.csv", library)
    assert result.checks.no_sample and result.next_step.startswith("CO2 artefact")


def test_saturation_is_flagged(library):
    result = example("06_PS_calibration_film_transmission.csv", library, "transmission")
    assert result.checks.saturated and "WARNING: bands saturated" in result.next_step


def test_analysis_never_changes_the_raw_data(library):
    spectrum = read_spectra(EXAMPLES / "06_PS_calibration_film_transmission.csv")[0]
    before = spectrum.raw.copy()
    analyse(spectrum, library, "transmission")
    assert np.array_equal(spectrum.raw, before)


# ---------------------------------------------------------------- ATR and transmission
KNOWN_MODES = {f.name: ("ATR" if "ATR" in f.name else "transmission") for f in sorted(EXAMPLES.glob("*.csv"))
               if "ATR" in f.name or "transmission" in f.name}


def as_transmission(spectrum):
    """The same measurement as it would look in transmission: ATR bands at high wavenumber are weaker."""
    a = spectrum.raw * spectrum.wn / config.REFERENCE_WN
    return replace(spectrum, raw=a, unit="A")


def test_a_transmission_sample_meets_the_atr_library(library):
    spectrum = read_spectra(EXAMPLES / "01_wipe_area_particle_ATR.csv")[0]
    by_atr = analyse(spectrum, library, "ATR")
    by_transmission = analyse(as_transmission(spectrum), library, "transmission")
    assert by_transmission.r == pytest.approx(by_atr.r, abs=1e-9) and by_transmission.match == by_atr.match == "PET"


def test_an_atr_sample_meets_a_transmission_reference(tmp_path):
    from ftir_analyzer.inhouse import InHouseLibrary
    spectrum = read_spectra(EXAMPLES / "12_shipping_box_film_ATR.csv")[0]
    film = as_transmission(spectrum)
    InHouseLibrary(tmp_path).add("EVOH film, transmission", film.wn, film.raw, "transmission")
    result = analyse(spectrum, load_library(inhouse_dir=tmp_path), "ATR")
    assert result.best == "EVOH film, transmission" and result.r == pytest.approx(1.0)
    InHouseLibrary(tmp_path / "wrong").add("EVOH film, transmission", film.wn, film.raw, "ATR")   # mode set wrong
    wrong = analyse(spectrum, load_library(inhouse_dir=tmp_path / "wrong"), "ATR")
    assert dict(wrong.ranking)["EVOH film, transmission"] < 0.9               # why the reference's mode matters


@pytest.mark.parametrize("name", sorted(KNOWN_MODES))
def test_an_unknown_mode_is_worked_out_from_the_better_fit(library, name):
    known, unknown = example(name, library, KNOWN_MODES[name]), example(name, library, "unknown")
    assert (unknown.verdict, unknown.match) == (known.verdict, known.match)
    other_r = unknown.other[1]
    if abs(unknown.r - other_r) > 0.01:                                        # not a tie (demo 10: 0.876 / 0.879)
        assert unknown.reading == KNOWN_MODES[name]
    assert ("mode unknown", "warn") in unknown.flags and "fits better as" in unknown.next_step
    assert ("mode doubt", "warn") not in known.flags                           # the right mode is never doubted


def test_a_wrong_mode_is_noticed(library):
    result = example("02_wafer_box_particle_ATR.csv", library, "transmission")
    assert ("mode doubt", "warn") in result.flags and "fits clearly better as ATR" in result.next_step
    result = example("06_PS_calibration_film_transmission.csv", library, "ATR")
    assert ("mode doubt", "warn") in result.flags


def test_no_reading_fits_so_nothing_is_corrected(library):
    result = example("08_fibre_on_filter.csv", library)
    assert result.reading == "unknown" and result.compared_as == "as measured" and result.verdict == NO_MATCH


# ---------------------------------------------------------------- records
def test_mode_is_asked_once_and_saved(tmp_path):
    log = SampleLog(tmp_path / "sample_log.csv")
    noted, unnoted = EXAMPLES / "01_wipe_area_particle_ATR.csv", EXAMPLES / "07_particle_on_filter.csv"
    asked = []

    def ask(paths):
        asked.extend(paths)
        return {p: "ATR" for p in paths}

    first = resolve_modes([noted, unnoted], log, ask)
    assert asked == [unnoted]                          # the file note answers for the first file
    assert first[noted].source == "file note" and first[unnoted].mode == "ATR"
    second = resolve_modes([unnoted], log, ask)
    assert asked == [unnoted] and second[unnoted].source.startswith("sample log")   # not asked again


def test_dont_know_is_not_saved(tmp_path):
    log = SampleLog(tmp_path / "sample_log.csv")
    path = EXAMPLES / "08_fibre_on_filter.csv"
    assert resolve_modes([path], log, lambda paths: {p: "unknown" for p in paths})[path].mode == "unknown"
    assert log.entries() == {}


def test_report_and_append_only_audit_log(tmp_path, library):
    results = [example("01_wipe_area_particle_ATR.csv", library, "ATR")]
    info = run_info(library, impostor_test(library))
    audit = tmp_path / "audit_log.csv"
    for run in ("run1", "run2"):
        html_path = write_report(results, library, info, tmp_path / run, audit_log=audit)
        assert html_path.exists() and (tmp_path / run / "report.csv").exists()
    assert len(pd.read_csv(audit)) == 2                # the second run was added, the first kept
    assert "Ali Obeidat" in html_path.read_text(encoding="utf-8")


def test_a_signed_report_and_its_record(tmp_path, library):
    from ftir_analyzer.records import make_signoff, readable_time
    from ftir_analyzer.spectrum_io import checksum
    results = [example("05_fibre_on_wafer_ATR.csv", library, "ATR")]
    info = run_info(library, impostor_test(library), "Ali Obeidat")
    with pytest.raises(ValueError):
        make_signoff("", "accept")                                    # a signature needs a name
    signoff = make_signoff("Ali Obeidat", "accept with changes", "05 is Nylon 6, see the band at 1640")
    signatures = tmp_path / "signatures.csv"
    html_path = write_report(results, library, info, tmp_path / "signed", audit_log=tmp_path / "audit.csv",
                             signoff=signoff, signature_log=signatures)
    page = html_path.read_text(encoding="utf-8")
    seen = re.sub(r"<[^>]+>", "", page)                               # the text as the reader sees it
    assert "Signed: Accepted with changes" in seen and "Ali Obeidat" in page and "1 spectrum " in page
    assert readable_time(info["run_at"], with_zone=True) in page and info["run_at"] not in page.split("<main>")[0]
    row = pd.read_csv(signatures, comment="#").iloc[0]
    assert row.signed_by == "Ali Obeidat" and row.decision == "accept with changes"
    assert row.report_sha256 == checksum(html_path)                   # the signature points at this exact file
    unsigned = write_report(results, library, info, tmp_path / "unsigned", audit_log=tmp_path / "audit.csv",
                            signature_log=signatures)
    assert "Not signed yet" in re.sub(r"<[^>]+>", "", unsigned.read_text(encoding="utf-8")) and len(pd.read_csv(signatures, comment="#")) == 1


def test_pdf_report_has_a_page_per_spectrum(tmp_path, library):
    from ftir_analyzer.records import make_signoff
    names = ("01_wipe_area_particle_ATR.csv", "07_particle_on_filter.csv", "08_fibre_on_filter.csv")   # clean, second band, no match
    results = [example(name, library, "ATR") for name in names]
    info = run_info(library, impostor_test(library), "Ali Obeidat")
    for run, signoff in (("open", None), ("signed", make_signoff("Ali Obeidat", "accept", "checked by hand"))):
        html_path = write_report(results, library, info, tmp_path / run, audit_log=tmp_path / "audit.csv",
                                 signoff=signoff, signature_log=tmp_path / "signatures.csv")
        pdf = (html_path.parent / "report.pdf").read_bytes()
        assert pdf.startswith(b"%PDF")
        assert len(re.findall(rb"/Type /Page(?!s)", pdf)) >= 1 + len(results) + 1    # summary, one per spectrum, traceability


def test_bands_in_the_report_can_be_pointed_at(tmp_path, library):
    results = [example("01_wipe_area_particle_ATR.csv", library, "ATR")]
    page = write_report(results, library, run_info(library, impostor_test(library)), tmp_path / "run",
                        audit_log=tmp_path / "audit.csv").read_text(encoding="utf-8")
    assert 'class="chip band"' in page and 'data-fig="result0"' in page and 'class="band-card" hidden' in page
    assert "bandmark" in page                                         # the script that marks a band in the figure
    assert 'id="motion-toggle"' in page and "window.ftirMotion" in page   # the motion switch in the bar
    assert "Rules used" in page and 'class="grid2 even"' in page          # the traceability: two panels of one height


def test_readable_times():
    from ftir_analyzer.records import readable_time
    assert readable_time("2026-10-20T16:50:00+02:00") == "20 Oct 2026, 16:50"
    assert readable_time("2026-10-20T16:50:00+02:00", with_zone=True) == "20 Oct 2026, 16:50 (UTC+02:00)"
    assert readable_time("2026-10-20") == "20 Oct 2026" and readable_time("not a time") == "not a time"


def test_records_open_as_formatted_excel_sheets(tmp_path, library):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from ftir_analyzer.log_views import FIRST_ROW, view_of
    results = [example(name, library, "ATR") for name in ("01_wipe_area_particle_ATR.csv", "08_fibre_on_filter.csv")]
    audit = tmp_path / "audit_log.csv"
    write_report(results, library, run_info(library, impostor_test(library)), tmp_path / "run", audit_log=audit)
    before = audit.read_bytes()
    view = view_of(audit, "Audit log", tmp_path / "views")
    assert audit.read_bytes() == before                               # the record itself is never touched
    ws = load_workbook(view)["Audit log"]
    assert ws["A1"].value == "Audit log" and "Rows are only ever added" in ws["A2"].value
    assert ws.cell(FIRST_ROW, 1).value == "Run at" and ws.freeze_panes == f"A{FIRST_ROW + 1}"
    assert ws.cell(FIRST_ROW + 1, 1).is_date                          # a real date, not text
    verdicts = [ws.cell(FIRST_ROW + 1 + i, [c.value for c in ws[FIRST_ROW]].index("Verdict") + 1) for i in range(2)]
    assert [v.value for v in verdicts] == ["probable", "no reliable match"]
    assert verdicts[0].fill.fgColor.rgb.endswith("E5F3EA")            # the verdict in its colour
    assert load_workbook(tmp_path / "run" / "report.xlsx").sheetnames == ["Results", "About"]   # the report's table


# ---------------------------------------------------------------- analytical view and CO2 direction
def test_analytical_view_reproduces_every_step(library):
    from ftir_analyzer.analytical import analytical_steps
    result = example("06_PS_calibration_film_transmission.csv", library, "transmission")
    steps = analytical_steps(result, library)
    assert [s.number for s in steps] == [1, 2, 3, 4, 5]
    raw, absorbance = steps[0].curves[0], steps[1].curves[0]
    assert np.allclose(absorbance.y, transmittance_to_absorbance(raw.y))      # step 2 = -log10(T)
    after = steps[2].curves[1]
    assert np.array_equal(after.y, result.compared, equal_nan=True)            # step 3 = what was compared
    sample_fp, reference_fp = steps[3].curves
    assert pearson(sample_fp.y, reference_fp.y) == pytest.approx(result.r)     # step 4 = the score
    assert {name for name, _ in steps[4].rulers} == {"C–H", "O–H / N–H", "CO₂"}


def test_co2_artefact_points_down_in_the_empty_spot(library):
    checks = example("09_empty_spot_on_filter.csv", library).checks
    assert checks.co2_direction == "down"          # less CO2 in the air than during the background scan
    assert checks.co2_signed < 0 and checks.below_zero


def test_interactive_figures(library):
    pytest.importorskip("plotly")
    from ftir_analyzer.interactive import analytical_figure, result_figure
    result = example("07_particle_on_filter.csv", library, "ATR")
    figure = result_figure(result, library)
    assert {"raw absorbance", "sample, as compared (ATR)"} <= {trace.name for trace in figure.data}
    assert "y5" in [trace.yaxis for trace in analytical_figure(result, library).data]   # five panels
    # a fixed wavenumber range, 4000 on the left: autorange broke when the window changed size
    assert figure.layout.xaxis.range == (4000.0, 650.0) and figure.layout.xaxis.autorange is False


def test_report_page_works_on_every_screen_and_on_paper(tmp_path, library):
    pytest.importorskip("plotly")
    from ftir_analyzer.report import figure_colour_map
    swap = figure_colour_map()                                 # every figure colour has exactly one dark partner
    assert swap["#FFFFFF"] == "#1A1918" and len(set(swap.values())) == len(swap)
    results = [example(name, library, "ATR") for name in ("06_PS_calibration_film_transmission.csv",
                                                          "07_particle_on_filter.csv")]
    page = write_report(results, library, run_info(library, impostor_test(library)), tmp_path / "run",
                        audit_log=tmp_path / "audit.csv").read_text(encoding="utf-8")
    assert "@media screen and (prefers-color-scheme:dark)" in page   # dark colours never reach the paper
    assert "beforeprint" in page and "FIGURE_SWAP" in page           # print opens every closed block
    assert "06_<wbr>PS_<wbr>calibration" in page                     # long names wrap at the underscores
    assert "font-family:'IBM Plex Mono'" in page                     # checksums look the same everywhere
    assert "weak from </em>0.60" in page and "probable from </em>0.95" in page   # the score strip
    assert "&#10003;" not in page                                    # icons, not text symbols


def test_figure_labels_never_touch(library):
    """Every label of the result view, on every demo spectrum, at a wide and a narrow size: no two
    labels overlap, and 'unexplained' never sits on its own dashed line."""
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.text import Annotation
    from ftir_analyzer.plots import result_figure
    for path in sorted(EXAMPLES.glob("*.csv")):
        result = example(path.name, library, "ATR")
        for size in ((10, 5.2), (6, 4)):
            fig = result_figure(result, library, size=size)
            FigureCanvasAgg(fig)
            fig._fit_labels()
            renderer = fig.canvas.get_renderer()
            bottom = fig.axes[1]
            texts = [a for a in bottom.texts if a.get_visible() and a.get_text()] + [fig.axes[0].title, bottom.title]
            boxes = [a.get_window_extent(renderer) for a in texts] + [fig.legends[0].get_window_extent(renderer)]
            for i, a in enumerate(boxes):
                for b in boxes[i + 1:]:
                    assert not a.overlaps(b), f"{path.name} {size}: two labels overlap"
            lines = [ln for ln in bottom.get_lines() if len(ln.get_xydata()) == 2]   # the dashed 'unexplained' lines
            for text in [a for a in bottom.texts if a.get_visible() and a.get_text()]:
                box = text.get_window_extent(renderer)
                for ln in lines:
                    (x, y0), (_, y1) = ln.get_transform().transform(ln.get_xydata())
                    crosses = box.x0 < x < box.x1 and min(y0, y1) < box.y1 and max(y0, y1) > box.y0
                    assert not crosses, f"{path.name} {size}: '{text.get_text()}' crosses a line"
            for note in [a for a in bottom.texts if isinstance(a, Annotation) and "unexplained" in a.get_text()]:
                line_x = bottom.transData.transform((note.xy[0], 0))[0]
                box = note.get_window_extent(renderer)
                assert not box.x0 - 2 <= line_x <= box.x1 + 2, f"{path.name}: label on its line"


def test_water_vapour_lines(library):
    """The polystyrene film (demo 06, a real measurement) has the sharp comb of water-vapour lines;
    the clean demos do not; data too coarse to show the lines is 'not measured', never 'clean'."""
    film = example("06_PS_calibration_film_transmission.csv", library, "transmission")
    assert film.checks.vapour_lines and ("water-vapour lines", "warn") in film.flags
    assert "purge longer" in film.next_step
    for name in ("01_wipe_area_particle_ATR.csv", "07_particle_on_filter.csv", "12_shipping_box_film_ATR.csv"):
        assert not example(name, library, "ATR").checks.vapour_lines
    coarse = example("08_fibre_on_filter.csv", library, "ATR")           # a point every 7.7 cm-1
    assert np.isnan(coarse.checks.vapour) and "not measured" in describe_vapour(coarse.checks)


def test_two_component_fit(library):
    """A made-up mix of two references is recognised with the right partner and about the right share;
    a pure spectrum and a band no reference has are left alone (no invented mix)."""
    from ftir_analyzer.mixture import two_component_fit
    from ftir_analyzer.spectrum_io import Spectrum
    rng = np.random.default_rng(1)

    def scaled(y):
        return (y - np.nanmin(y)) / np.nanmax(y - np.nanmin(y))
    for first, second, share in (("PE (high density)", "PET", 0.3), ("PP", "PMMA", 0.25)):
        mix = (1 - share) * scaled(library.spectra[first]) + share * scaled(library.spectra[second])
        mix = mix + rng.normal(0, 0.003, len(GRID))
        keep = np.isfinite(mix)
        result = analyse(Spectrum("mix", GRID[keep].copy(), mix[keep].copy(), "A", "test", EXAMPLES, "-"),
                         library, "ATR")
        assert result.best == first and result.evidence.unexplained          # the second material shows
        assert result.mixture.other == second
        assert result.mixture.share_other == pytest.approx(share, abs=0.05)
        assert result.mixture.r_fit > result.mixture.r_alone + config.MIX_GAIN
    assert example("01_wipe_area_particle_ATR.csv", library, "ATR").mixture is None     # pure PET
    particle = example("07_particle_on_filter.csv", library, "ATR")                    # 1030: in no reference
    assert particle.evidence.unexplained and two_component_fit(particle, library) is None


def test_report_motion_is_built_in_and_optional(tmp_path, library):
    """GSAP and Lenis are inside the page (it works offline), the motion steps aside for 'reduce
    motion' and for printing, and the traceability keeps the impostor test as one line."""
    page = write_report([example("07_particle_on_filter.csv", library, "ATR")], library,
                        run_info(library, impostor_test(library)), tmp_path / "run",
                        audit_log=tmp_path / "audit.csv").read_text(encoding="utf-8")
    assert "GSAP 3.15.0" in page and "1.3.27" in page and "<script src" not in page   # nothing from the internet
    assert "prefers-reduced-motion: reduce" in page and "window.ftirMotion" in page   # the page steps aside, or follows the switch
    assert "beforeprint" in page and "SplitText.create" in page and "ScrollTrigger" in page
    assert "Validation" in page and "Validation and data integrity" not in page
